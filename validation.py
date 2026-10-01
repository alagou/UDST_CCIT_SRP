"""Validation issues for the SRP pipeline.

Issues are dicts with Severity, Project_ID, Field, Issue, and Raw_Value.
They are written to ``output/SRP_Validation_Report.xlsx``.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ISSUE_COLUMNS = ["Severity", "Project_ID", "Field", "Issue", "Raw_Value"]
_SEVERITY_ORDER = {"ERROR": 0, "WARNING": 1, "INFO": 2}


def add_issue(issues: list, severity: str, project_id, field: str, issue: str, raw_value) -> None:
    """Append one validation finding. Raw values are shortened for Excel."""
    text = "" if raw_value is None else str(raw_value)
    if len(text) > 500:
        text = text[:497] + "..."
    issues.append(
        {
            "Severity": severity,
            "Project_ID": project_id or "",
            "Field": field or "",
            "Issue": issue,
            "Raw_Value": text,
        }
    )


def add_cross_checks(projects: list, memberships: list, registry_ids: list[str], issues: list, config: dict) -> None:
    """Add dataset-level checks that a single cell cleaner cannot see."""
    for project_id in registry_ids:
        add_issue(
            issues,
            "ERROR",
            project_id,
            "Project_ID",
            "Duplicate Project_ID in the registry",
            project_id,
        )

    seen_ids = Counter(project["project_id"] for project in projects)
    for project_id, count in seen_ids.items():
        if count > 1:
            add_issue(
                issues,
                "ERROR",
                project_id,
                "Project_ID",
                "Duplicate Project_ID in the project table",
                count,
            )

    for project in projects:
        _check_project(project, issues)

    _check_membership_rows(projects, memberships, issues, config)


def _check_project(project: dict, issues: list) -> None:
    project_id = project["project_id"]
    if not project.get("supervisor_name"):
        add_issue(issues, "ERROR", project_id, "Supervisor", "Missing supervisor name", project.get("raw_supervisor_name"))
    email = project.get("supervisor_email")
    email_warning = project.get("email_warning")
    if not email:
        add_issue(issues, "ERROR", project_id, "Supervisor_Email", "Missing supervisor email", project.get("raw_supervisor_email"))
    elif email_warning == "malformed email":
        add_issue(
            issues,
            "ERROR",
            project_id,
            "Supervisor_Email",
            "Malformed supervisor email",
            project.get("raw_supervisor_email"),
        )
    elif email_warning == "numeric local-part":
        add_issue(
            issues,
            "WARNING",
            project_id,
            "Supervisor_Email",
            "Supervisor email looks like a numeric ID",
            project.get("raw_supervisor_email"),
        )
    if not project.get("title"):
        add_issue(issues, "ERROR", project_id, "Project Title", "Missing project title", project.get("raw_title"))

    maximum = project.get("max_students")
    capacity_warning = project.get("capacity_warning")
    raw_capacity = project.get("raw_capacity")
    if maximum is None:
        add_issue(issues, "ERROR", project_id, "Maximum Number of Students", "Missing or non-numeric capacity", raw_capacity)
    elif maximum <= 0:
        add_issue(issues, "ERROR", project_id, "Maximum Number of Students", "Invalid capacity", raw_capacity)
    elif capacity_warning:
        add_issue(
            issues,
            "WARNING",
            project_id,
            "Maximum Number of Students",
            "Non-numeric capacity; interpreted as the largest number found",
            raw_capacity,
        )
    assigned = project.get("assigned_count") or 0
    if maximum is not None and maximum > 0 and assigned > maximum:
        add_issue(
            issues,
            "ERROR",
            project_id,
            "Maximum Number of Students",
            "Assigned students exceed stated capacity",
            f"{assigned} assigned, maximum {maximum}",
        )
    if project.get("year_warning"):
        add_issue(
            issues,
            "WARNING",
            project_id,
            "Preferred Student Year & Program",
            "Unparseable preferred year/program",
            project.get("raw_year_program"),
        )
    if project.get("name_id_mismatch"):
        add_issue(
            issues,
            "WARNING",
            project_id,
            "Student ID(s)",
            "name/ID count mismatch",
            project.get("name_id_mismatch"),
        )
    # Admin labels replace the computed label. Record that skip every time.
    if project.get("status_source") == "admin-override":
        add_issue(
            issues,
            "INFO",
            project_id,
            "Status",
            "Admin status override applied; computed recruitment logic was skipped",
            f"admin={project.get('status_label')}; computed={project.get('computed_label')}",
        )
        if project.get("status_label") != project.get("computed_label"):
            add_issue(
                issues,
                "WARNING",
                project_id,
                "Status",
                "Conflicting status values between Admin Form and computed logic",
                f"admin={project.get('status_label')}; computed={project.get('computed_label')}",
            )


def _check_membership_rows(projects: list, memberships: list, issues: list, config: dict) -> None:
    for project in projects:
        if project.get("name_id_mismatch"):
            continue
        for member in project.get("members") or []:
            _check_member_pair(project["project_id"], member, issues)
        counts = Counter(
            member["student_id"]
            for member in project.get("members") or []
            if member.get("id_valid") and member.get("membership_status") == "Assigned"
        )
        for student_id, count in counts.items():
            if count > 1:
                add_issue(
                    issues,
                    "WARNING",
                    project["project_id"],
                    "Student_ID",
                    "Duplicate Student_ID on the same project",
                    student_id,
                )

    if not config.get("membership_exclusive", True):
        return
    by_student: dict[str, set] = {}
    for member in memberships:
        if not member.get("id_valid") or member.get("membership_status") != "Assigned":
            continue
        by_student.setdefault(member["student_id"], set()).add(member["project_id"])
    for student_id, project_ids in sorted(by_student.items()):
        if len(project_ids) > 1:
            add_issue(
                issues,
                "WARNING",
                ", ".join(sorted(project_ids)),
                "Student_ID",
                "Duplicate Student_ID across different projects",
                student_id,
            )


def _check_member_pair(project_id: str, member: dict, issues: list) -> None:
    name = member.get("student_name")
    student_id = member.get("student_id")
    if name and not student_id:
        add_issue(issues, "WARNING", project_id, "Student_ID", "Student name present without a Student_ID", name)
    if student_id and not name:
        add_issue(issues, "WARNING", project_id, "Student Name", "Student_ID present without a student name", student_id)
    if student_id and not member.get("id_valid"):
        add_issue(issues, "WARNING", project_id, "Student_ID", "Invalid Student_ID", member.get("raw_student_id") or student_id)


def count_by_severity(issues: list) -> dict[str, int]:
    counts = {"ERROR": 0, "WARNING": 0, "INFO": 0}
    for issue in issues:
        severity = issue.get("Severity", "")
        if severity in counts:
            counts[severity] += 1
    return counts


def write_validation_report(path: Path, issues: list) -> None:
    """Write one sortable row per issue."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(issues)
    if not rows:
        rows = [
            {
                "Severity": "INFO",
                "Project_ID": "",
                "Field": "",
                "Issue": "No validation issues.",
                "Raw_Value": "",
            }
        ]
    rows.sort(key=lambda row: (_SEVERITY_ORDER.get(row["Severity"], 9), row["Project_ID"], row["Field"], row["Issue"]))
    frame = pd.DataFrame(rows, columns=ISSUE_COLUMNS)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Validation", index=False)
        _style_sheet(writer.sheets["Validation"], {"A": 12, "B": 28, "C": 36, "D": 78, "E": 55})


def _style_sheet(worksheet, widths: dict) -> None:
    header_fill = PatternFill("solid", fgColor="0F2744")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
